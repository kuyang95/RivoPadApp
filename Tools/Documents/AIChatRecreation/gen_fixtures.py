"""Generate ExcelAIChatInternetRecreationTests.swift from downloaded workbooks.

Each fixture recreates one sheet at the ORIGINAL cell addresses using only the
app's AI chat: table creation (anchored at the selected cell), parameter
mini-tables, chunked row fills (<=100 edits per turn), number formats,
dropdowns, and conditional formatting.
"""
import datetime, json, re, sys, warnings
import openpyxl
warnings.filterwarnings("ignore")

NET = "/private/tmp/claude-501/-Users-me-Develop-IOSProject-RivoPadApp/1acd2f02-1101-4b02-ab57-575301ff9796/scratchpad/net/"
OUT = "/Users/me/Develop/IOSProject/RivoPadApp/shortcuts_exampleTests/ExcelAIChatInternetRecreationTests.swift"

def col_letter(c):
    s = ""
    while c > 0:
        c, r = divmod(c - 1, 26)
        s = chr(65 + r) + s
    return s

def norm(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return repr(v)
    return str(v)

def swift_str(s):
    return '"' + s.replace("\\", "\\\\").replace('"', '\\"') + '"'

def swift_multiline(s):
    # Triple-quoted literal; escape backslashes and triple quotes only.
    body = s.replace("\\", "\\\\").replace('"""', '\\"\\"\\"')
    lines = body.split("\n")
    return '"""\n' + "\n".join("                    " + l for l in lines) + '\n                    """'

class Fixture:
    def __init__(self, name, output):
        self.name = name
        self.output = output
        self.turns = []  # (select, request, verify_swift)

    def turn(self, request, verify, select=None):
        self.turns.append((select, request, verify))

    def verify_only(self, verify):
        self.turns.append((None, None, verify))

def read_block(ws, header_row, first_col, last_col, first_data_row, last_data_row):
    headers = [norm(ws.cell(header_row, c).value) for c in range(first_col, last_col + 1)]
    rows = []
    for r in range(first_data_row, last_data_row + 1):
        rows.append([norm(ws.cell(r, c).value) for c in range(first_col, last_col + 1)])
    return headers, rows

def chunk(rows, size):
    for i in range(0, len(rows), size):
        yield i, rows[i:i + size]

def table_turns(fx, ws, header_row, first_col, last_col, first_data_row, last_data_row,
                fill_chunk, extra_fill_hint=""):
    headers, rows = read_block(ws, header_row, first_col, last_col, first_data_row, last_data_row)
    anchor = f"{col_letter(first_col)}{header_row}"
    end = f"{col_letter(last_col)}{last_data_row}"
    n = len(rows)
    fx.turn(
        f"선택한 셀 {anchor}부터 {', '.join(headers)} 열을 이 순서로 가진 정식 Excel 표를 만들어줘. 빈 데이터 행은 정확히 {n}개로 해줘.",
        f'try expectTable(workbook, range: "{anchor}:{end}", headers: {json.dumps(headers, ensure_ascii=False)})',
        select=anchor,
    )
    for offset, part in chunk(rows, fill_chunk):
        r0 = first_data_row + offset
        r1 = r0 + len(part) - 1
        lines = "\n".join(f"{r0 + i}행: " + " | ".join(v for v in row) for i, row in enumerate(part))
        hint = (" 각 줄 맨 앞의 'N행:'은 그 줄이 들어갈 행 번호야. 값이 비어 있는 칸은 빈 셀로 두고, "
                "=로 시작하는 값은 수식으로 그대로 넣어줘. 적혀 있지 않은 내용은 만들지 마. "
                "표 끝에 새 행을 추가하지 말고 표 안에 이미 있는 그 행 번호의 빈 셀을 채워야 해." + extra_fill_hint)
        fx.turn(
            f"{anchor}:{end} 표의 {r0}행부터 {r1}행까지 아래 내용으로 채워줘. 각 줄은 {' | '.join(headers)} 순서야.{hint}\n{lines}",
            f'try expectValues(workbook, rows: {json.dumps(part, ensure_ascii=False)}, startRow: {r0}, startColumn: {first_col})',
        )
    return headers, rows

def explicit_cells_turn(fx, cells):
    """cells: list of (reference, text). One chat turn naming each address."""
    parts = [f"{ref} 셀에 '{text}'" for ref, text in cells]
    verify = "\n                        ".join(
        f'try expectValues(workbook, rows: [[{swift_str(text)}]], startRow: {int(re.sub("[A-Z]", "", ref))}, startColumn: {col_index(re.sub("[0-9]", "", ref))})'
        for ref, text in cells
    )
    fx.turn(
        "표 밖의 셀에 값을 직접 써줘. " + ", ".join(parts) + " 를 그대로 넣어줘. 다른 셀은 건드리지 마.",
        verify,
    )

def col_index(letters):
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n

def param_table_turns(fx, header, value, header_cell, value_row, col, number_format=None, format_name=None):
    anchor = header_cell
    fx.turn(
        f"선택한 셀 {anchor}부터 {header} 열 하나만 있는 정식 Excel 표를 만들어줘. 빈 데이터 행은 정확히 1개로 해줘.",
        f'try expectTable(workbook, range: "{anchor}:{col_letter(col)}{value_row}", headers: [{swift_str(header)}])',
        select=anchor,
    )
    fx.turn(
        f"{anchor}:{col_letter(col)}{value_row} 표의 {value_row}행 {header} 값을 {value}로 넣어줘.",
        f'try expectValues(workbook, rows: [[{swift_str(value)}]], startRow: {value_row}, startColumn: {col})',
    )
    if number_format:
        fx.turn(
            f"{anchor}:{col_letter(col)}{value_row} 표의 {header} 열을 {number_format} 형식으로 표시해줘.",
            f'try expectNumberFormat(workbook, .{format_name}, column: {col}, rows: {value_row}...{value_row})',
        )

# ---------------------------------------------------------------- fixtures
fixtures = []

# 1) ExcelDemy Supermarket Sales -------------------------------------------
ws = openpyxl.load_workbook(NET + "ed-supermarket.xlsx")["SuperMarket Sales"]
fx = Fixture("Supermarket Sales (ExcelDemy)", "AIChat_Net_SupermarketSales.xlsx")
explicit_cells_turn(fx, [("B2", "Excel Sample Data"), ("B4", "Supermarket Sales Data"), ("H6", "Tax"), ("I6", "0.1")])
headers, rows = table_turns(fx, ws, 8, 2, 9, 9, 80, fill_chunk=10)
fx.verify_only(
    "try expectNoValues(workbook, rows: 81...120, columns: 1...12)\n"
    "                        try expectNumberFormat(workbook, .date, column: 3, rows: 9...80)\n"
    "                        try expectNumberFormat(workbook, .date, column: 5, rows: 9...80)\n"
    f"                        try expectValues(workbook, rows: {json.dumps(rows, ensure_ascii=False)}, startRow: 9, startColumn: 2)",
)
fixtures.append(fx)

# 2) ExcelDemy Students Marksheet ---------------------------------------------
ws = openpyxl.load_workbook(NET + "ed-marksheet.xlsx")["Student Marks"]
fx = Fixture("Students Marksheet (ExcelDemy)", "AIChat_Net_StudentsMarksheet.xlsx")
explicit_cells_turn(fx, [("B2", "Excel Sample Data"), ("B4", "Students Marksheet Data"), ("F6", "Total Marks"), ("G6", "300")])
headers, rows = table_turns(fx, ws, 8, 2, 7, 9, 38, fill_chunk=15)
fx.turn(
    "B8:G38 표의 Percentage 열을 백분율로 표시해줘.",
    "try expectNoValues(workbook, rows: 39...80, columns: 1...10)\n"
    "                        try expectNumberFormat(workbook, .percent, column: 7, rows: 9...38)\n"
    f"                        try expectValues(workbook, rows: {json.dumps(rows, ensure_ascii=False)}, startRow: 9, startColumn: 2)\n"
    '                        let g9 = try sheet(workbook).cell(at: ExcelCellAddress(row: 9, column: 7))\n'
    '                        guard g9?.displayValue == "84%" else { throw fail("G9 표시값 \\(g9?.displayValue ?? "") != 84%") }',
)
fixtures.append(fx)

# 3) Smartsheet Project Tracking ----------------------------------------------
ws = openpyxl.load_workbook(NET + "ss-project-tracking.xlsx")["Project Tracking"]
fx = Fixture("Project Tracking (Smartsheet)", "AIChat_Net_ProjectTracking.xlsx")
explicit_cells_turn(fx, [("B2", "PROJECT TRACKING TEMPLATE"), ("F3", "PROJECTS"), ("I3", "DELIVERABLE(S)"), ("K3", "COST / HOURS"), ("B31", "CLICK HERE TO CREATE IN SMARTSHEET")])
# key tables (header text without the embedded newline)
fx.turn(
    "선택한 셀 O4부터 STATUS KEY 열 하나만 있는 정식 Excel 표를 만들어줘. 빈 데이터 행은 정확히 6개로 해줘.",
    'try expectTable(workbook, range: "O4:O10", headers: ["STATUS KEY"])',
    select="O4",
)
status_keys = ["Not Started", "In Progress", "Complete", "On Hold", "Overdue"]
fx.turn(
    "O4:O10 표의 5행부터 9행까지 STATUS KEY 값을 순서대로 채워줘.\n" + "\n".join(status_keys),
    f'try expectValues(workbook, rows: {json.dumps([[k] for k in status_keys])}, startRow: 5, startColumn: 15)',
)
fx.turn(
    "선택한 셀 Q4부터 PRIORITY KEY 열 하나만 있는 정식 Excel 표를 만들어줘. 빈 데이터 행은 정확히 4개로 해줘.",
    'try expectTable(workbook, range: "Q4:Q8", headers: ["PRIORITY KEY"])',
    select="Q4",
)
priority_keys = ["High", "Medium", "Low"]
fx.turn(
    "Q4:Q8 표의 5행부터 7행까지 PRIORITY KEY 값을 순서대로 채워줘.\n" + "\n".join(priority_keys),
    f'try expectValues(workbook, rows: {json.dumps([[k] for k in priority_keys])}, startRow: 5, startColumn: 17)',
)
headers, rows = table_turns(fx, ws, 4, 2, 13, 5, 29, fill_chunk=8,
                            extra_fill_hint=" 비어 있는 줄도 행 번호를 지키기 위해 그대로 두어야 하니 건너뛰지 말고 빈 값으로 처리해줘.")
fx.turn(
    "B4:M29 표의 STATUS 열 데이터 셀에 Not Started, In Progress, Complete, On Hold, Overdue 중에서 고르는 드롭다운을 넣고, PRIORITY 열에는 High, Medium, Low 드롭다운을 넣어줘.",
    'try expectDropdown(workbook, values: ["Not Started", "In Progress", "Complete", "On Hold", "Overdue"], column: 3, rows: 5...29)\n'
    '                        try expectDropdown(workbook, values: ["High", "Medium", "Low"], column: 4, rows: 5...29)',
)
status_cf = [("Overdue", "red"), ("On Hold", "yellow"), ("Complete", "green"), ("In Progress", "yellow"), ("Not Started", "yellow")]
priority_cf = [("Low", "green"), ("Medium", "yellow"), ("High", "red")]
color_ko = {"red": "빨간색", "yellow": "노란색", "green": "초록색"}
fx.turn(
    "B4:M29 표에 '텍스트 포함' 조건부 서식을 걸어줘. 값이 정확히 같을 때가 아니라 그 글자를 포함하면 되는 조건이야. STATUS 열은 "
    + ", ".join(f"'{t}' 를 포함하는 셀은 {color_ko[c]}" for t, c in status_cf)
    + ". PRIORITY 열은 "
    + ", ".join(f"'{t}' 를 포함하는 셀은 {color_ko[c]}" for t, c in priority_cf)
    + ". 규칙 8개를 전부 각각 만들어줘.",
    "\n                        ".join(
        f'try expectConditionalRule(workbook, kind: .containsText, comparisonValue: "{t}", highlight: .{c}, column: 3, rows: 5...29)' for t, c in status_cf
    ) + "\n                        " + "\n                        ".join(
        f'try expectConditionalRule(workbook, kind: .containsText, comparisonValue: "{t}", highlight: .{c}, column: 4, rows: 5...29)' for t, c in priority_cf
    ),
)
fx.turn(
    "B4:M29 표의 % DONE 열을 백분율로 표시해줘.",
    "try expectNoValues(workbook, rows: 32...80, columns: 1...20)\n"
    "                        try expectNumberFormat(workbook, .percent, column: 10, rows: 5...29)\n"
    f"                        try expectValues(workbook, rows: {json.dumps(rows, ensure_ascii=False)}, startRow: 5, startColumn: 2)",
)
fixtures.append(fx)

# ---------------------------------------------------------------- emit Swift
def emit_turn(select, request, verify):
    sel = f"select: {swift_str(select)}, " if select else ""
    req = swift_multiline(request) if request is not None else "nil"
    return f"""                Turn(
                    {sel}request: {req},
                    verify: {{ workbook in
                        {verify}
                    }}
                ),"""

fixture_code = []
for fx in fixtures:
    turns = "\n".join(emit_turn(*t) for t in fx.turns)
    fixture_code.append(f"""        // {fx.name}
        Fixture(
            outputName: {swift_str(fx.output)},
            turns: [
{turns}
            ]
        ),""")

HARNESS = open(NET + "harness.swift.txt").read()
swift = HARNESS.replace("__FIXTURES__", "\n".join(fixture_code))
open(OUT, "w").write(swift)
print("wrote", OUT, "turns:", [len(f.turns) for f in fixtures])
