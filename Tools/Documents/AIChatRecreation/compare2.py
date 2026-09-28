"""Compare recreated workbooks with the downloaded originals at identical addresses."""
import sys, datetime, warnings
import openpyxl
warnings.filterwarnings("ignore")
NET = "/private/tmp/claude-501/-Users-me-Develop-IOSProject-RivoPadApp/1acd2f02-1101-4b02-ab57-575301ff9796/scratchpad/net/"
CASES = [
    ("AIChat_Net_SupermarketSales.xlsx", "ed-supermarket.xlsx", "SuperMarket Sales", []),
    ("AIChat_Net_StudentsMarksheet.xlsx", "ed-marksheet.xlsx", "Student Marks", []),
    ("AIChat_Net_ProjectTracking.xlsx", "ss-project-tracking.xlsx", "Project Tracking", []),
]

def norm(v):
    if v is None: return ""
    if isinstance(v, datetime.datetime): return v.strftime("%Y-%m-%d")
    if isinstance(v, bool): return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip().replace("\n", " ")

def main(out_dir):
    for out_name, orig_name, sheet, skipped_titles in CASES:
        print(f"\n===== {out_name}  <=  {orig_name} / '{sheet}'")
        wo = openpyxl.load_workbook(NET + orig_name)[sheet]
        try:
            wn = openpyxl.load_workbook(f"{out_dir}/{out_name}").worksheets[0]
        except Exception as e:
            print("  LOAD FAIL:", e); continue
        total = mism = 0
        details = []
        fmt_pairs = {}
        max_row = min(wo.max_row, 200); max_col = min(wo.max_column, 30)
        for r in range(1, max_row + 1):
            for c in range(1, max_col + 1):
                o = wo.cell(r, c)
                if o.value is None: continue
                if o.coordinate in skipped_titles: continue
                n = wn.cell(r, c)
                total += 1
                ov, nv = norm(o.value), norm(n.value)
                if ov != nv:
                    try:
                        if float(ov) == float(nv): continue
                    except Exception: pass
                    mism += 1; details.append((o.coordinate, ov, nv))
                if not str(o.value).startswith("="):
                    key = (o.number_format, n.number_format)
                    if key[0] != key[1]: fmt_pairs[key] = fmt_pairs.get(key, 0) + 1
        extra = [(cell.coordinate, norm(cell.value)) for row in wn.iter_rows() for cell in row
                 if cell.value is not None and wo.cell(cell.row, cell.column).value is None]
        print(f"  original non-empty cells compared: {total}, mismatches: {mism}, skipped title cells: {skipped_titles}")
        for d in details[:12]: print("   ", d)
        print("  cells we have that the original lacks:", extra[:12], "..." if len(extra) > 12 else "")
        print("  number-format differences (original -> ours: count):")
        for (a, b), k in sorted(fmt_pairs.items(), key=lambda kv: -kv[1]): print(f"    {a!r} -> {b!r}: {k}")
        print("  original DV:", [(str(d.sqref), d.formula1) for d in wo.data_validations.dataValidation])
        print("  ours     DV:", [(str(d.sqref), d.formula1) for d in wn.data_validations.dataValidation])
        print("  original CF:", [(str(r.sqref), [(x.type, x.operator, (x.formula or [''])[0][:40]) for x in r.rules]) for r in wo.conditional_formatting])
        print("  ours     CF:", [(str(r.sqref), [(x.type, x.operator, (x.formula or [''])[0][:40]) for x in r.rules]) for r in wn.conditional_formatting])
        print("  original tables:", list(wo.tables.items()), " ours:", list(wn.tables.items()))

if __name__ == "__main__":
    main(sys.argv[1])
